# ─────────────────────────────────────────────────────────────────────────────
# backend/parsers/docx_parser.py
# DOCX and PPTX parsing using python-docx / python-pptx directly.
# ─────────────────────────────────────────────────────────────────────────────
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


async def parse_docx(file_bytes: bytes, filename: str) -> str:
    """Parse DOCX or PPTX file into plain text."""
    suffix = Path(filename).suffix.lower()
    if suffix in (".pptx", ".ppt"):
        return await _parse_pptx(file_bytes, filename)
    return await _parse_docx(file_bytes, filename)


async def _parse_docx(file_bytes: bytes, filename: str) -> str:
    """Parse DOCX using python-docx."""
    import io
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        text = "\n\n".join(paragraphs)
        logger.info(f"[DOCX Parser] Extracted {len(text)} chars from '{filename}'")
        return text
    except Exception as e:
        logger.error(f"[DOCX Parser] Failed: {e}")
        raise ValueError(f"Could not parse DOCX '{filename}': {e}")


async def _parse_pptx(file_bytes: bytes, filename: str) -> str:
    """Parse PPTX using python-pptx."""
    import io
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(file_bytes))
        slides_text = []
        for i, slide in enumerate(prs.slides, 1):
            slide_parts = [f"--- Slide {i} ---"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
            slides_text.append("\n".join(slide_parts))
        text = "\n\n".join(slides_text)
        logger.info(f"[PPTX Parser] Extracted {len(text)} chars from '{filename}'")
        return text
    except Exception as e:
        logger.error(f"[PPTX Parser] Failed: {e}")
        raise ValueError(f"Could not parse PPTX '{filename}': {e}")
